import docx
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor
import numpy as np
import pandas as pd

# ==========================================
# 1. FILE MAPPING & DATA EXTRACTION SETUP
# ==========================================

FILES = {
    'M4': 'M4_Fully_Coupled_SSI_Baseline.xlsx',
    'M1': 'M1_Uniform_Winkler.xlsx',
    'M2a': 'M2a_Pseudo_Coduto.xlsx',
    'M2b': 'M2b_Pseudo_Bowles.xlsx',
    'M2c': 'M2c_Pseudo_Hanse_George.xlsx',
    'M3': 'M3_Semi_Coupled_SSI.xlsx',
}

MODELS = ['M4', 'M1', 'M2a', 'M2b', 'M2c', 'M3']
MODEL_LABELS = {
    'M4': 'M4 (Baseline)',
    'M1': 'M1 (Uniform)',
    'M2a': 'M2a (Coduto)',
    'M2b': 'M2b (Bowles)',
    'M2c': 'M2c (Hansen-G.)',
    'M3': 'M3 (Semi-Coupled)',
}


def get_val(file_path, sheet_name, x, y, possible_cols):
    """Extracts scalar values from target cutline sheet with sign conventions applied[cite: 1]."""
    df = pd.read_excel(file_path, sheet_name=sheet_name)
    sub = df[(abs(df['X_m'] - x) < 1e-3) & (abs(df['Y_m'] - y) < 1e-3)]
    if sub.empty:
        df['dist'] = (df['X_m'] - x) ** 2 + (df['Y_m'] - y) ** 2
        sub = df.loc[[df['dist'].idxmin()]]

    val = None
    for col in possible_cols:
        if col in sub.columns:
            val = sub[col].values[0]
            break
    if val is None:
        numeric_cols = [
            c
            for c in sub.select_dtypes(include=[np.number]).columns
            if c not in ['X_m', 'Y_m']
        ]
        val = sub[numeric_cols[0]].values[0]

    if pd.isna(val):
        return 0.0

    val = float(val)

    # Sign convention adjustments for consistency across ETABS/PLAXIS outputs
    if file_path in [FILES['M4'], FILES['M3']]:
        if any(m in col_name for col_name in possible_cols for m in ['M11', 'M22']):
            val = -val

    if file_path in [FILES['M1'], FILES['M2a'], FILES['M2b'], FILES['M2c']]:
        if any(
            p in col_name
            for col_name in possible_cols
            for p in ['Effective_Soil_Stress_kPa', 'Soil_Pressure_kPa']
        ):
            val = -val

    return val


def calc_delta(val, baseline_val):
    """Calculates percentage discrepancy relative to M4 baseline[cite: 1]."""
    if baseline_val == 0:
        return '—'
    if (val * baseline_val) < 0:
        return 'N/A'
    delta = ((val - baseline_val) / baseline_val) * 100.0
    if abs(delta) < 0.05:
        return '0.0'
    return f'{delta:+.1f}'


# ==========================================
# 2. WORD DOCUMENT STYLING & XML HELPERS
# ==========================================


def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)


def set_cell_margins(cell, top=100, bottom=100, start=120, end=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [
        ('top', top),
        ('bottom', bottom),
        ('left', start),
        ('right', end),
    ]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def set_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f"""
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="12" w:space="0" w:color="1F497D"/>
            <w:bottom w:val="single" w:sz="12" w:space="0" w:color="1F497D"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="D3D3D3"/>
            <w:insideV w:val="none"/>
            <w:left w:val="none"/>
            <w:right w:val="none"/>
        </w:tblBorders>
    """)
    tblPr.append(borders)


def format_run(run, font_name='Calibri', size_pt=9, bold=False, color_rgb=None):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    if color_rgb:
        run.font.color.rgb = color_rgb


# ==========================================
# 3. TABLE BUILDER FUNCTIONS
# ==========================================


def create_geotech_table(
    doc, title, footnote, col_headers, sheet, p1, p2, p3, p4, p5
):
    """Generates 11-column Geotechnical Table (Uz and Pressure)[cite: 1]."""
    p = doc.add_paragraph()
    r = p.add_run(title)
    format_run(r, size_pt=11, bold=True, color_rgb=RGBColor(31, 73, 125))
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(2)

    if footnote:
        p_fn = doc.add_paragraph()
        r_fn = p_fn.add_run(f'Footnote: {footnote}')
        format_run(r_fn, size_pt=8.5, color_rgb=RGBColor(80, 80, 80))
        p_fn.paragraph_format.space_after = Pt(4)

    table = doc.add_table(rows=2 + len(MODELS), cols=11)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)

    # Header Rows
    hdr0, hdr1 = table.rows[0], table.rows[1]
    hdr0.cells[0].merge(hdr1.cells[0])
    hdr0.cells[0].paragraphs[0].text = 'Model Approach'

    cats = [
        (1, col_headers[0], col_headers[1]),
        (3, col_headers[2], col_headers[3]),
        (5, col_headers[4], col_headers[5]),
        (7, col_headers[6], col_headers[7]),
        (9, col_headers[8], col_headers[9]),
    ]

    for idx, (col_idx, label, unit) in enumerate(cats):
        cell = hdr0.cells[col_idx]
        cell.merge(hdr0.cells[col_idx + 1])
        cell.paragraphs[0].text = label
        hdr1.cells[col_idx].paragraphs[0].text = f'Value ({unit})'
        hdr1.cells[col_idx + 1].paragraphs[0].text = 'δ (%)'

    for r_idx in range(2):
        for cell in table.rows[r_idx].cells:
            set_cell_background(cell, '1F497D')
            set_cell_margins(cell, top=120, bottom=120)
            for cp in cell.paragraphs:
                cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for cr in cp.runs:
                    format_run(
                        cr, size_pt=9, bold=True, color_rgb=RGBColor(255, 255, 255)
                    )

    # Populate Data
    for r_i, m_key in enumerate(MODELS):
        row_cells = table.rows[2 + r_i].cells
        bg_color = 'F2F4F8' if r_i % 2 == 1 else 'FFFFFF'

        row_cells[0].paragraphs[0].text = MODEL_LABELS[m_key]
        row_cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
        row_cells[0].paragraphs[0].runs[0].font.bold = True

        file_path = FILES[m_key]
        v1 = get_val(file_path, sheet, p1['x'], p1['y'], p1['cols'])
        v2 = get_val(file_path, sheet, p2['x'], p2['y'], p2['cols'])
        v_diff = v1 - v2
        v3 = get_val(file_path, sheet, p4['x'], p4['y'], p4['press_cols'])
        v4 = get_val(file_path, sheet, p5['x'], p5['y'], p5['press_cols'])

        m4_v1 = get_val(FILES['M4'], sheet, p1['x'], p1['y'], p1['cols'])
        m4_v2 = get_val(FILES['M4'], sheet, p2['x'], p2['y'], p2['cols'])
        m4_diff = m4_v1 - m4_v2
        m4_v3 = get_val(FILES['M4'], sheet, p4['x'], p4['y'], p4['press_cols'])
        m4_v4 = get_val(FILES['M4'], sheet, p5['x'], p5['y'], p5['press_cols'])

        row_data_parsed = [
            (v1, calc_delta(v1, m4_v1)),
            (v2, calc_delta(v2, m4_v2)),
            (v_diff, calc_delta(v_diff, m4_diff)),
            (v3, calc_delta(v3, m4_v3)),
            (v4, calc_delta(v4, m4_v4)),
        ]

        col_cursor = 1
        for val, d_str in row_data_parsed:
            c_val = row_cells[col_cursor]
            set_cell_background(c_val, bg_color)
            set_cell_margins(c_val, top=80, bottom=80)
            p_c = c_val.paragraphs[0]
            r_c = p_c.add_run(f'{val:.2f}')
            format_run(r_c, size_pt=9)
            p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER

            c_d = row_cells[col_cursor + 1]
            set_cell_background(c_d, bg_color)
            set_cell_margins(c_d, top=80, bottom=80)
            p_d = c_d.paragraphs[0]
            r_d = p_d.add_run(d_str if d_str in ['—', 'N/A'] else f'{d_str}%')
            format_run(r_d, size_pt=9)
            p_d.alignment = WD_ALIGN_PARAGRAPH.CENTER

            col_cursor += 2

    doc.add_paragraph().paragraph_format.space_after = Pt(8)


def create_struct_table(
    doc, title, footnote, col_headers, sheet, zone_a, zone_b1, zone_b2, zone_c, mom_col
):
    """Generates 9-column Structural Bending Moment Table[cite: 1]."""
    p = doc.add_paragraph()
    r = p.add_run(title)
    format_run(r, size_pt=11, bold=True, color_rgb=RGBColor(31, 73, 125))
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(2)

    if footnote:
        p_fn = doc.add_paragraph()
        r_fn = p_fn.add_run(f'Footnote: {footnote}')
        format_run(r_fn, size_pt=8.5, color_rgb=RGBColor(80, 80, 80))
        p_fn.paragraph_format.space_after = Pt(4)

    table = doc.add_table(rows=2 + len(MODELS), cols=9)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)

    # Header Rows
    hdr0, hdr1 = table.rows[0], table.rows[1]
    hdr0.cells[0].merge(hdr1.cells[0])
    hdr0.cells[0].paragraphs[0].text = 'Model Approach'

    zones = [
        (1, col_headers[0], zone_a),
        (3, col_headers[1], zone_b1),
        (5, col_headers[2], zone_b2),
        (7, col_headers[3], zone_c),
    ]

    for start_col, label, z_pt in zones:
        cell = hdr0.cells[start_col]
        cell.merge(hdr0.cells[start_col + 1])
        cell.paragraphs[0].text = label
        hdr1.cells[start_col].paragraphs[0].text = 'kNm/m'
        hdr1.cells[start_col + 1].paragraphs[0].text = 'δ (%)'

    for r_idx in range(2):
        for cell in table.rows[r_idx].cells:
            set_cell_background(cell, '2F5597')
            set_cell_margins(cell, top=120, bottom=120)
            for cp in cell.paragraphs:
                cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for cr in cp.runs:
                    format_run(
                        cr, size_pt=9, bold=True, color_rgb=RGBColor(255, 255, 255)
                    )

    m4_a = get_val(FILES['M4'], sheet, zone_a['x'], zone_a['y'], [mom_col])
    m4_b1 = get_val(FILES['M4'], sheet, zone_b1['x'], zone_b1['y'], [mom_col])
    m4_b2 = get_val(FILES['M4'], sheet, zone_b2['x'], zone_b2['y'], [mom_col])
    m4_c = get_val(FILES['M4'], sheet, zone_c['x'], zone_c['y'], [mom_col])

    # Populate Data
    for r_i, m_key in enumerate(MODELS):
        row_cells = table.rows[2 + r_i].cells
        bg_color = 'F2F4F8' if r_i % 2 == 1 else 'FFFFFF'

        row_cells[0].paragraphs[0].text = MODEL_LABELS[m_key]
        row_cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
        row_cells[0].paragraphs[0].runs[0].font.bold = True

        f_path = FILES[m_key]
        va = get_val(f_path, sheet, zone_a['x'], zone_a['y'], [mom_col])
        vb1 = get_val(f_path, sheet, zone_b1['x'], zone_b1['y'], [mom_col])
        vb2 = get_val(f_path, sheet, zone_b2['x'], zone_b2['y'], [mom_col])
        vc = get_val(f_path, sheet, zone_c['x'], zone_c['y'], [mom_col])

        data_struct = [
            (va, calc_delta(va, m4_a)),
            (vb1, calc_delta(vb1, m4_b1)),
            (vb2, calc_delta(vb2, m4_b2)),
            (vc, calc_delta(vc, m4_c)),
        ]

        col_cursor = 1
        for val, d_str in data_struct:
            c_val = row_cells[col_cursor]
            set_cell_background(c_val, bg_color)
            set_cell_margins(c_val, top=80, bottom=80)
            p_c = c_val.paragraphs[0]
            r_c = p_c.add_run(f'{val:.2f}')
            format_run(r_c, size_pt=9)
            p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER

            c_d = row_cells[col_cursor + 1]
            set_cell_background(c_d, bg_color)
            set_cell_margins(c_d, top=80, bottom=80)
            p_d = c_d.paragraphs[0]
            r_d = p_d.add_run(d_str if d_str in ['—', 'N/A'] else f'{d_str}%')
            format_run(r_d, size_pt=9)
            p_d.alignment = WD_ALIGN_PARAGRAPH.CENTER

            col_cursor += 2

    doc.add_paragraph().paragraph_format.space_after = Pt(8)


# ==========================================
# 4. MASTER DOCUMENT BUILDER
# ==========================================


def build_master_document():
    doc = docx.Document()

    # Set Landscape with 0.75-inch margins
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11.0)
    section.page_height = Inches(8.5)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    # Document Header Title
    title_p = doc.add_paragraph()
    title_r = title_p.add_run(
        'CHAPTER 4: COMPARATIVE ANALYSIS TABLES ACROSS 7 CUTLINE AXES'
    )
    format_run(
        title_r, size_pt=16, bold=True, color_rgb=RGBColor(31, 73, 125)
    )
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_after = Pt(18)

    sett_cols = ['Uz_mm', 'Settlement_mm']
    press_cols = ['Effective_Soil_Stress_kPa', 'Soil_Pressure_kPa']

    # ------------------------------------------------------------------------
    # SECTION 1: Cut_Y=10m (Longitudinal Centerline Axis)
    # ------------------------------------------------------------------------
    sec1_p = doc.add_paragraph()
    sec1_r = sec1_p.add_run('1. Longitudinal Centerline Axis (Cut_Y = 10 m)')
    format_run(sec1_r, size_pt=13, bold=True, color_rgb=RGBColor(31, 73, 125))

    create_geotech_table(
        doc,
        title='Table 1A: Geotechnical Response Parameters along Cut_Y = 10 m',
        footnote='Evaluated across the longitudinal centerline from exterior soil overhang (X = -1 m) to structural center (X = 15 m).',
        col_headers=[
            'Center Sett. Uz,center at (15,10)',
            'mm',
            'Edge Sett. Uz,edge at (-1,10)',
            'mm',
            'Diff. Sett. ΔUz (-1–15 m)',
            'mm',
            'Center Press. pcenter at (15,10)',
            'kPa',
            'Edge Press. pedge at (-1,10)',
            'kPa',
        ],
        sheet='Cut_Y=10m',
        p1={'x': 15.0, 'y': 10.0, 'cols': sett_cols},
        p2={'x': -1.0, 'y': 10.0, 'cols': sett_cols},
        p3={'x': 15.0, 'y': 10.0, 'cols': sett_cols},
        p4={'x': 15.0, 'y': 10.0, 'press_cols': press_cols},
        p5={'x': -1.0, 'y': 10.0, 'press_cols': press_cols},
    )

    create_struct_table(
        doc,
        title='Table 1B: Structural Bending Moment M11 along Cut_Y = 10 m',
        footnote='Evaluated across structural zones: Zone A (Interior Column), Zone B1 (Mid-Span), Zone B2 (Edge Span), Zone C (Perimeter Edge).',
        col_headers=[
            'Zone A: Column Line (15,10)',
            'Zone B1: Mid-Span (7.5,10)',
            'Zone B2: Edge span (2.5,10)',
            'Zone C: Perimeter Edge (-1,10)',
        ],
        sheet='Cut_Y=10m',
        zone_a={'x': 15.0, 'y': 10.0},
        zone_b1={'x': 7.5, 'y': 10.0},
        zone_b2={'x': 2.5, 'y': 10.0},
        zone_c={'x': -1.0, 'y': 10.0},
        mom_col='M11_kNm_m',
    )

    # ------------------------------------------------------------------------
    # SECTION 2: Cut_X=15m (Transverse Centerline Axis)
    # ------------------------------------------------------------------------
    sec2_p = doc.add_paragraph()
    sec2_r = sec2_p.add_run('2. Transverse Centerline Axis (Cut_X = 15 m)')
    format_run(sec2_r, size_pt=13, bold=True, color_rgb=RGBColor(31, 73, 125))

    create_geotech_table(
        doc,
        title='Table 2A: Geotechnical Response Parameters along Cut_X = 15 m',
        footnote='Evaluated across the transverse centerline from exterior soil overhang (Y = -1 m) to structural center (Y = 10 m).',
        col_headers=[
            'Center Sett. Uz,center at (15,10)',
            'mm',
            'Edge Sett. Uz,edge at (15,-1)',
            'mm',
            'Diff. Sett. ΔUz (-1–10 m)',
            'mm',
            'Center Press. pcenter at (15,10)',
            'kPa',
            'Edge Press. pedge at (15,-1)',
            'kPa',
        ],
        sheet='Cut_X=15m',
        p1={'x': 15.0, 'y': 10.0, 'cols': sett_cols},
        p2={'x': 15.0, 'y': -1.0, 'cols': sett_cols},
        p3={'x': 15.0, 'y': 10.0, 'cols': sett_cols},
        p4={'x': 15.0, 'y': 10.0, 'press_cols': press_cols},
        p5={'x': 15.0, 'y': -1.0, 'press_cols': press_cols},
    )

    create_struct_table(
        doc,
        title='Table 2B: Structural Bending Moment M22 along Cut_X = 15 m',
        footnote='Evaluated across structural zones: Zone A (Interior Column), Zone B1 (Mid-Span), Zone B2 (Edge Span), Zone C (Perimeter Edge).',
        col_headers=[
            'Zone A: Column Line (15,10)',
            'Zone B1: Mid-Span (15,7.5)',
            'Zone B2: Edge span (15,2.5)',
            'Zone C: Perimeter Edge (15,-1)',
        ],
        sheet='Cut_X=15m',
        zone_a={'x': 15.0, 'y': 10.0},
        zone_b1={'x': 15.0, 'y': 7.5},
        zone_b2={'x': 15.0, 'y': 2.5},
        zone_c={'x': 15.0, 'y': -1.0},
        mom_col='M22_kNm_m',
    )

    # ------------------------------------------------------------------------
    # SECTION 3: Cut_Y=0m (Longitudinal Boundary Edge Axis)
    # ------------------------------------------------------------------------
    sec3_p = doc.add_paragraph()
    sec3_r = sec3_p.add_run(
        '3. Longitudinal Boundary Edge Axis (Cut_Y = 0 m)'
    )
    format_run(sec3_r, size_pt=13, bold=True, color_rgb=RGBColor(31, 73, 125))

    create_geotech_table(
        doc,
        title='Table 3A: Geotechnical Response Parameters along Cut_Y = 0 m',
        footnote='Evaluated along the boundary edge from exterior overhang (X = -1 m) to structural center (X = 15 m).',
        col_headers=[
            'Interior Sett. Uz,interior at (15,0)',
            'mm',
            'Edge Sett. Uz,edge at (-1,0)',
            'mm',
            'Diff. Sett. ΔUz (-1–15 m)',
            'mm',
            'Interior Press. pinterior at (15,0)',
            'kPa',
            'Edge Press. pedge at (-1,0)',
            'kPa',
        ],
        sheet='Cut_Y=0m',
        p1={'x': 15.0, 'y': 0.0, 'cols': sett_cols},
        p2={'x': -1.0, 'y': 0.0, 'cols': sett_cols},
        p3={'x': 15.0, 'y': 0.0, 'cols': sett_cols},
        p4={'x': 15.0, 'y': 0.0, 'press_cols': press_cols},
        p5={'x': -1.0, 'y': 0.0, 'press_cols': press_cols},
    )

    create_struct_table(
        doc,
        title='Table 3B: Structural Bending Moment M11 along Cut_Y = 0 m',
        footnote='Evaluated across boundary structural zones along Y = 0 m.',
        col_headers=[
            'Zone A: Column Line (15,0)',
            'Zone B1: Mid-Span (7.5,0)',
            'Zone B2: Edge span (2.5,0)',
            'Zone C: Perimeter Edge (-1,0)',
        ],
        sheet='Cut_Y=0m',
        zone_a={'x': 15.0, 'y': 0.0},
        zone_b1={'x': 7.5, 'y': 0.0},
        zone_b2={'x': 2.5, 'y': 0.0},
        zone_c={'x': -1.0, 'y': 0.0},
        mom_col='M11_kNm_m',
    )

    # ------------------------------------------------------------------------
    # SECTION 4: Cut_X=0m (Transverse Boundary Edge Axis)
    # ------------------------------------------------------------------------
    sec4_p = doc.add_paragraph()
    sec4_r = sec4_p.add_run('4. Transverse Boundary Edge Axis (Cut_X = 0 m)')
    format_run(sec4_r, size_pt=13, bold=True, color_rgb=RGBColor(31, 73, 125))

    create_geotech_table(
        doc,
        title='Table 4A: Geotechnical Response Parameters along Cut_X = 0 m',
        footnote='Evaluated along the boundary edge from exterior overhang (Y = -1 m) to structural center (Y = 10 m).',
        col_headers=[
            'Interior Sett. Uz,interior at (0,10)',
            'mm',
            'Edge Sett. Uz,edge at (0,-1)',
            'mm',
            'Diff. Sett. ΔUz (-1–10 m)',
            'mm',
            'Interior Press. pinterior at (0,10)',
            'kPa',
            'Edge Press. pedge at (0,-1)',
            'kPa',
        ],
        sheet='Cut_X=0m',
        p1={'x': 0.0, 'y': 10.0, 'cols': sett_cols},
        p2={'x': 0.0, 'y': -1.0, 'cols': sett_cols},
        p3={'x': 0.0, 'y': 10.0, 'cols': sett_cols},
        p4={'x': 0.0, 'y': 10.0, 'press_cols': press_cols},
        p5={'x': 0.0, 'y': -1.0, 'press_cols': press_cols},
    )

    create_struct_table(
        doc,
        title='Table 4B: Structural Bending Moment M22 along Cut_X = 0 m',
        footnote='Evaluated across boundary structural zones along X = 0 m.',
        col_headers=[
            'Zone A: Column Line (0,10)',
            'Zone B1: Mid-Span (0,7.5)',
            'Zone B2: Edge span (0,2.5)',
            'Zone C: Perimeter Edge (0,-1)',
        ],
        sheet='Cut_X=0m',
        zone_a={'x': 0.0, 'y': 10.0},
        zone_b1={'x': 0.0, 'y': 7.5},
        zone_b2={'x': 0.0, 'y': 2.5},
        zone_c={'x': 0.0, 'y': -1.0},
        mom_col='M22_kNm_m',
    )

    # ------------------------------------------------------------------------
    # SECTION 5: Cut_X=5m (Intermediate Transverse Axis X = 5 m)
    # ------------------------------------------------------------------------
    sec5_p = doc.add_paragraph()
    sec5_r = sec5_p.add_run('5. Intermediate Transverse Axis (Cut_X = 5 m)')
    format_run(sec5_r, size_pt=13, bold=True, color_rgb=RGBColor(31, 73, 125))

    create_geotech_table(
        doc,
        title='Table 5A: Geotechnical Response Parameters along Cut_X = 5 m',
        footnote='Evaluated along intermediate transverse axis X = 5 m from exterior overhang (Y = -1 m) to centerline (Y = 10 m).',
        col_headers=[
            'Interior Sett. Uz,interior at (5,10)',
            'mm',
            'Edge Sett. Uz,edge at (5,-1)',
            'mm',
            'Diff. Sett. ΔUz (-1–10 m)',
            'mm',
            'Interior Press. pinterior at (5,10)',
            'kPa',
            'Edge Press. pedge at (5,-1)',
            'kPa',
        ],
        sheet='Cut_X=5m',
        p1={'x': 5.0, 'y': 10.0, 'cols': sett_cols},
        p2={'x': 5.0, 'y': -1.0, 'cols': sett_cols},
        p3={'x': 5.0, 'y': 10.0, 'cols': sett_cols},
        p4={'x': 5.0, 'y': 10.0, 'press_cols': press_cols},
        p5={'x': 5.0, 'y': -1.0, 'press_cols': press_cols},
    )

    create_struct_table(
        doc,
        title='Table 5B: Structural Bending Moment M22 along Cut_X = 5 m',
        footnote='Evaluated across structural zones along intermediate axis X = 5 m.',
        col_headers=[
            'Zone A: Column Line (5,10)',
            'Zone B1: Mid-Span (5,7.5)',
            'Zone B2: Edge span (5,2.5)',
            'Zone C: Perimeter Edge (5,-1)',
        ],
        sheet='Cut_X=5m',
        zone_a={'x': 5.0, 'y': 10.0},
        zone_b1={'x': 5.0, 'y': 7.5},
        zone_b2={'x': 5.0, 'y': 2.5},
        zone_c={'x': 5.0, 'y': -1.0},
        mom_col='M22_kNm_m',
    )

    # ------------------------------------------------------------------------
    # SECTION 6: Cut_X=10m (Intermediate Transverse Axis X = 10 m)
    # ------------------------------------------------------------------------
    sec6_p = doc.add_paragraph()
    sec6_r = sec6_p.add_run('6. Intermediate Transverse Axis (Cut_X = 10 m)')
    format_run(sec6_r, size_pt=13, bold=True, color_rgb=RGBColor(31, 73, 125))

    create_geotech_table(
        doc,
        title='Table 6A: Geotechnical Response Parameters along Cut_X = 10 m',
        footnote='Evaluated along intermediate transverse axis X = 10 m from exterior overhang (Y = -1 m) to centerline (Y = 10 m).',
        col_headers=[
            'Interior Sett. Uz,interior at (10,10)',
            'mm',
            'Edge Sett. Uz,edge at (10,-1)',
            'mm',
            'Diff. Sett. ΔUz (-1–10 m)',
            'mm',
            'Interior Press. pinterior at (10,10)',
            'kPa',
            'Edge Press. pedge at (10,-1)',
            'kPa',
        ],
        sheet='Cut_X=10m',
        p1={'x': 10.0, 'y': 10.0, 'cols': sett_cols},
        p2={'x': 10.0, 'y': -1.0, 'cols': sett_cols},
        p3={'x': 10.0, 'y': 10.0, 'cols': sett_cols},
        p4={'x': 10.0, 'y': 10.0, 'press_cols': press_cols},
        p5={'x': 10.0, 'y': -1.0, 'press_cols': press_cols},
    )

    create_struct_table(
        doc,
        title='Table 6B: Structural Bending Moment M22 along Cut_X = 10 m',
        footnote='Evaluated across structural zones along intermediate axis X = 10 m.',
        col_headers=[
            'Zone A: Column Line (10,10)',
            'Zone B1: Mid-Span (10,7.5)',
            'Zone B2: Edge span (10,2.5)',
            'Zone C: Perimeter Edge (10,-1)',
        ],
        sheet='Cut_X=10m',
        zone_a={'x': 10.0, 'y': 10.0},
        zone_b1={'x': 10.0, 'y': 7.5},
        zone_b2={'x': 10.0, 'y': 2.5},
        zone_c={'x': 10.0, 'y': -1.0},
        mom_col='M22_kNm_m',
    )

    # ------------------------------------------------------------------------
    # SECTION 7: Cut_Y=5m (Intermediate Longitudinal Axis Y = 5 m)
    # ------------------------------------------------------------------------
    sec7_p = doc.add_paragraph()
    sec7_r = sec7_p.add_run('7. Intermediate Longitudinal Axis (Cut_Y = 5 m)')
    format_run(sec7_r, size_pt=13, bold=True, color_rgb=RGBColor(31, 73, 125))

    create_geotech_table(
        doc,
        title='Table 7A: Geotechnical Response Parameters along Cut_Y = 5 m',
        footnote='Evaluated along intermediate longitudinal axis Y = 5 m from exterior overhang (X = -1 m) to centerline (X = 15 m).',
        col_headers=[
            'Interior Sett. Uz,interior at (15,5)',
            'mm',
            'Edge Sett. Uz,edge at (-1,5)',
            'mm',
            'Diff. Sett. ΔUz (-1–15 m)',
            'mm',
            'Interior Press. pinterior at (15,5)',
            'kPa',
            'Edge Press. pedge at (-1,5)',
            'kPa',
        ],
        sheet='Cut_Y=5m',
        p1={'x': 15.0, 'y': 5.0, 'cols': sett_cols},
        p2={'x': -1.0, 'y': 5.0, 'cols': sett_cols},
        p3={'x': 15.0, 'y': 5.0, 'cols': sett_cols},
        p4={'x': 15.0, 'y': 5.0, 'press_cols': press_cols},
        p5={'x': -1.0, 'y': 5.0, 'press_cols': press_cols},
    )

    create_struct_table(
        doc,
        title='Table 7B: Structural Bending Moment M11 along Cut_Y = 5 m',
        footnote='Evaluated across structural zones along intermediate axis Y = 5 m.',
        col_headers=[
            'Zone A: Column Line (15,5)',
            'Zone B1: Mid-Span (7.5,5)',
            'Zone B2: Edge span (2.5,5)',
            'Zone C: Perimeter Edge (-1,5)',
        ],
        sheet='Cut_Y=5m',
        zone_a={'x': 15.0, 'y': 5.0},
        zone_b1={'x': 7.5, 'y': 5.0},
        zone_b2={'x': 2.5, 'y': 5.0},
        zone_c={'x': -1.0, 'y': 5.0},
        mom_col='M11_kNm_m',
    )

    # Output Document
    output_filename = 'Master_Data_Tables_Chapter4.docx'
    doc.save(output_filename)
    print(f'Successfully generated master tables document: {output_filename}')


if __name__ == '__main__':
    build_master_document()
